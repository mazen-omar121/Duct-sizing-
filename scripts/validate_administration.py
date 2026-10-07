"""Independent checks of delivered CAD/PDF/registers and unchanged design duties."""
import csv
import json
import math
from pathlib import Path

import ezdxf
import fitz
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'generated' / 'administration'
BASELINE = ROOT / 'projects' / 'administration_rev08' / 'cad-package' / 'schedules'


def records(path):
    with path.open() as f:
        return list(csv.DictReader(f))


def main():
    schedule = OUT / 'cad-package' / 'schedules'
    sections = records(schedule / 'Duct_sections.csv')
    assert sections == records(BASELINE / 'Duct_sections.csv'), 'Duct design inputs changed'
    terminals = records(schedule / 'Terminals.csv')
    assert terminals == records(BASELINE / 'Terminals.csv'), 'Terminal duties or locations changed'
    instruments = records(schedule / 'Instruments.csv')
    assert instruments == records(BASELINE / 'Instruments.csv'), 'Instrument requirements changed'
    index = records(schedule / 'Instrument_drawing_index.csv')
    assert len(index) == len(instruments) == 85
    assert {r['tag'] for r in index} == {r['tag'] for r in instruments}
    detailed = ezdxf.readfile(OUT / 'Administration_HVAC_Detailed_Layout.dxf')
    assert detailed.units == ezdxf.units.MM and not detailed.audit().has_errors
    assert len([n for n in detailed.layouts.names() if n.startswith('D') and n.endswith('_A1')]) == 18
    dimensions = [e for block in detailed.blocks for e in block.query('DIMENSION')]
    assert len(dimensions) == 11
    # Overall known plan extents and the internal main offsets are actual DIMENSION entities.
    lengths=sorted(round(e.get_measurement()) for e in dimensions)
    assert lengths == sorted([1800,1700,3500,13000,20000,8350,950,3900,3400,2900,19500]), lengths
    plan=detailed.blocks.get('PLAN_INSTRUMENTATION_D007')
    assert len([e for e in plan.query('TEXT') if e.dxf.text=='3 AI TO HPCP-01 / D-015']) == 9
    assert len([e for e in plan.query('TEXT') if e.dxf.text=='(-) OUTDOOR REF.']) == 9
    schematic = ezdxf.readfile(OUT / 'Administration_HVAC_Schematic.dxf')
    assert not schematic.audit().has_errors
    checked_refs=0
    with fitz.open(OUT / 'Administration_Detailed_HVAC_Duct_Flow_Diagram.pdf') as pdf:
        assert len(pdf) == 18
        texts=[p.get_text() for p in pdf]
        for p in pdf:
            assert abs(p.rect.width*25.4/72-841)<.01 and abs(p.rect.height*25.4/72-594)<.01
        for r in index:
            for reference in r['drawings'].split(' / '):
                number=int(reference)
                assert number != 16, 'Index-only reference does not prove a tagged diagram'
                assert r['tag'] in texts[number-1], (r['tag'],reference)
                checked_refs+=1
        assert '18' in texts[17]
        # All text must fit the sheet. Full sheets are then visually reviewed for local overlap.
        out_of_sheet=[]
        for number,page in enumerate(pdf,1):
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines',[]):
                    for span in line['spans']:
                        if not page.rect.contains(fitz.Rect(span['bbox'])):
                            out_of_sheet.append((number,span['text']))
        assert not out_of_sheet,out_of_sheet
    with fitz.open(OUT / 'Administration_Detailed_HVAC_Duct_Flow_Diagram_Monochrome.pdf') as pdf:
        assert len(pdf)==18
    with fitz.open(OUT / 'Administration_HVAC_Schematic.pdf') as pdf:
        assert len(pdf)==4
    wb=openpyxl.load_workbook(OUT/'Administration_HVAC_Drawing_Registers_Rev09.xlsx',data_only=True)
    assert len(wb.sheetnames)==15 and wb['Instruments'].max_row==86
    assert wb['Instrument_drawing_index'].max_row==86
    for row in list(wb['Sizing_check'].values)[1:]:
        flow,w,h,d=row[3:7]
        area=math.pi*(d/1000)**2/4 if d else w*h/1e6
        assert abs(flow/1000/area-row[8])<1e-9
    assert wb['Sizing_check'].max_row==115
    report={'revision':9,'detailed_sheets':18,'schematic_sheets':4,'monochrome_sheets':18,
            'instrument_tags':85,'instrument_sheet_references_checked':checked_refs,
            'native_dimensions':len(dimensions),'room_signal_trunks':9,'outdoor_reference_annotations':9,
            'duct_inputs':'unchanged from Rev08','terminal_inputs':'unchanged from Rev08',
            'instrument_requirements':'unchanged from Rev08','DXF_audits':'passed',
            'workbook_sheets':15,'velocities_recalculated':114,
            'native_AutoCAD_plot':'not run','DWG_export':'use AutoCAD Save As from DXF',
            'engineering_holds':['pressure / outdoor-air balance','levels and coordination',
                                 'fitting K / fan ESP / OEM selections','I&C settings and final wiring']}
    (OUT/'independent_validation.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report))


if __name__=='__main__':
    main()
