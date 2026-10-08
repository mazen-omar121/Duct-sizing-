"""Main duct and return: a separate, detailed common-main DX method package.

All distribution drawings are functional diagrams. No surveyed ceiling route,
equipment capacity or actual fan ESP is inferred from diagram dimensions.
"""
from pathlib import Path
from math import ceil,pi,log10,sqrt
import csv,json,argparse,hashlib,zipfile,shutil
from collections import defaultdict
import fitz,ezdxf
from shapely.geometry import LineString
from shapely.ops import unary_union
from reportlab.pdfgen import canvas
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from cad import Scene,table,pdf_render_scene,export_schematic,MM,COL

BASE=Path(__file__).resolve().parents[1]
ap=argparse.ArgumentParser();ap.add_argument('--output-dir',default=str(BASE.parents[1]/'generated/main_duct_and_return'))
args=ap.parse_args();OUT=Path(args.output_dir);OUT.mkdir(parents=True,exist_ok=True)
TITLE='Main duct and return';COUNT=14;SHEETS=[];COVERED=defaultdict(set);DX_COVERED=defaultdict(set)
def read(n):
    with (BASE/'inputs'/n).open() as f:return list(csv.DictReader(f))
DATA=json.loads((BASE/'inputs/Design_basis.json').read_text());ROOMS={r['id']:r for r in DATA['rooms']}
ORDER=['02','04','08','05','06','09','07','03','01'];RA_ORDER=[r for r in ORDER if ROOMS[r]['rn']]
INSTRUMENTS=read('Instruments.csv');IO=read('IO_points.csv');TERMS=read('Terminals.csv')
TOTAL_SA=sum(float(r['sa']) for r in ROOMS.values());TOTAL_RA=sum(float(r['ra']) for r in ROOMS.values())
MAKEUP=TOTAL_SA-TOTAL_RA
SHORT={'02':'ELECTRICAL','05':'OPERATOR','04':'MEETING','08':'TELECOM','06':'CORRIDOR','09':'TOILETS','07':'BED / REST','03':'KITCHEN','01':'CLEAN AGENT'}
def write(n,rows):
    with (OUT/n).open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def numeric(v):return None if v in ['TBC','',None] else float(v)
def hydro(q,w,h):
    a=w*h/1e6;v=q/1000/a;dh=2*w*h/(w+h)/1000;re=1.2*v*dh/1.81e-5;ff=.02
    for _ in range(30):ff=1/(-2*log10(.00009/(3.7*dh)+2.51/(re*sqrt(ff))))**2
    return v,ff/dh*.6*v*v
MAINS=[]
for air,order,field,target in [('SA',ORDER,'sa',4.8),('RA',RA_ORDER,'ra',3.5)]:
    for i,room in enumerate(order):
        q=sum(float(ROOMS[r][field]) for r in order[i:]);w=600 if air=='SA' else 700;depth=1300 if air=='SA' else 900;v,fr=hydro(q,w,depth)
        MAINS.append(dict(tag=f'{air}-M{i+1:02}',air_type=air,room_takeoff=room,flow_before_takeoff_l_s=q,
                          takeoff_l_s=ROOMS[room][field],flow_after_takeoff_l_s=q-ROOMS[room][field],
                          clear_width_mm=w,clear_depth_mm=depth,velocity_m_s=round(v,6),straight_friction_pa_m=round(fr,6),
                          installed_length_m='TBC - surveyed route required',status='Functional sequential room-flow accounting only; actual grouped roof routes and 31 segment duties are in Roof_routes_Rev12 / MDR-012'))
write('Common_main_sizing.csv',MAINS)
DX=[]
for c in 'ABC':
    for pre,parameter,signal,action in [('PSH','DX discharge high pressure','OEM LOCAL','Hardwired compressor cutout; OEM reset / threshold'),
        ('PSL','DX suction low pressure','OEM LOCAL','OEM low-pressure protection; refrigerant / operating envelope settings'),
        ('TSLL','Evaporator low temperature / frost protection','OEM LOCAL','OEM anti-freeze protection; fan / compressor response'),
        ('LSH','Condensate drain-pan overflow','DI proposal','OEM compressor inhibit / alarm; maintain drain access'),
        ('PT-SUC','Suction refrigerant pressure','AI proposal','Monitor only if selected OEM permits pressure transmitter / factory data'),
        ('PT-LIQ','Liquid refrigerant pressure','AI proposal','Monitor only if selected OEM permits pressure transmitter / factory data'),
        ('OL-CMP','Compressor overload protection','OEM LOCAL','Local motor protection; expose approved fault status')]:
        DX.append(dict(tag=f'{pre}-DX-{c}',parameter=parameter,location='PAU-'+c,signal=signal,action=action,
                       range_or_setpoint='OEM refrigerant / pressure-temperature envelope / electrical ratings TBC',
                       status='NEW DX proposal, separate from 85 retained instruments; factory controller may expose approved contacts or communications'))
DX.append(dict(tag='DPT-SA-M01',parameter='Common supply duct static pressure',location='Common SA main; selected representative station',
               signal='AI proposal',action='Optional static control only if OEM fan/VFD is suitable; maintain room scheduled minimum flow',
               range_or_setpoint='0..1000 Pa screening range; operating setpoint / station TBC',status='NEW DX option; does not override fire or OEM safeties'))
write('Additional_DX_proposals.csv',DX)

def frame(n,title,scale='NTS'):
    s=Scene();s.rect(10,10,821,574,'INK',None,.45);s.text(18,23,TITLE.upper()+' - '+title,4.1,'INK',True)
    s.line([(18,30),(823,30)],'INK',.3);s.line([(10,544),(831,544)],'INK',.4)
    for x in [245,628,737]:s.line([(x,544),(x,584)],'INK',.25)
    s.text(18,553,'ADMINISTRATION HVAC / DX PAU OPTION',3.1,'INK',True);s.text(18,564,'COMMON SUPPLY MAIN + COMMON RETURN MAIN',2.7)
    s.text(18,578,'FOR ENGINEERING COORDINATION - NOT FOR CONSTRUCTION',2.3,'PENDING')
    s.text(253,553,'DRAWING TITLE',2.2,'DIM');s.text(253,563,title,3.2,'INK',True)
    s.text(253,578,'ROUTES / OEM SELECTION / CHECKED AND APPROVED: PENDING',2.3,'DIM')
    s.text(637,553,'ADM-HVAC-MDR-'+f'{n:03}',3,'INK',True);s.text(637,577,'SCALE: '+scale+' AT A1',2.5)
    s.text(745,552,'METHOD REVISION 01',2.3);s.text(745,562,'08 OCT 2026',2.5);s.text(745,578,f'SHEET {n} OF {COUNT}',2.4)
    return s
def heading(s,x,y,t,w):s.text(x,y,t,3.1,'INK',True);s.line([(x,y+4),(x+w,y+4)],'INK',.25)
def note(s,x,y,t,w=390,z=2.8):return s.para(x,y,t,w,z)+6
def footer(s,t):s.para(18,530,t,805,2.45,'DIM')
def sensor(s,x,y,tag,n,label=None,dx=False,z=2.55):
    s.circle(x,y,4.5,'INST',.3,'#FFFFFF');s.text(x,y+1.3,label or tag.split('-')[0],2.1,'INST',True,'center')
    s.text(x,y-8,tag,z,'INST',False,'center');(DX_COVERED if dx else COVERED)[tag].add(n)
def network(s,paths,air,width=7):
    polygon=unary_union([LineString(p).buffer(w/2,cap_style=2,join_style=2) for p,w in paths])
    boundary=polygon.boundary;lines=list(boundary.geoms) if hasattr(boundary,'geoms') else [boundary]
    for line in lines:s.line(list(line.coords),air,.5)
    return polygon
def pair_arrow(s,a,b,air):s.arrow(a,b,air,.28,2)
def blank_rows(n,title,rows,headers,ratios=None):
    s=frame(n,title);table(s,18,48,805,headers,rows,ratios,2.6,15);return s

# MDR-001: two-duty PAU bank, common roof headers and 15 room drops.
s=frame(1,'WHOLE SYSTEM AIRFLOW METHOD')
s.text(132,56,'AIR: GREEN SUPPLY / BLUE RETURN / ORANGE EXTRACT / TEAL OUTDOOR',2.65,'INK',True)
s.text(132,69,'MAGENTA: INSTRUMENT / REFRIGERANT AS LABELLED; DASHED: CONTROL SIGNAL',2.5,'INST')
s.rect(20,69,93,377,'DIM',None,.3);s.text(66,83,'DX PAU BANK',3.3,'INK',True,'center')
for c,y in zip('ABC',[133,245,357]):
    s.rect(31,y,68,55,'INK',None,.45);s.text(65,y+16,'PAU-'+c,3.3,'INK',True,'center')
    s.text(65,y+30,'50% unit',2.8,align='center');s.text(65,y+44,'MDR-004 / 005',2.5,align='center')
s.text(66,429,'2 DUTY + 1 STANDBY',2.6,'INK',True,'center')
sa_paths=[([(115,120),(755,120)],8)];ra_paths=[([(115,420),(691,420)],8)]
positions={r:(132+i*74) for i,r in enumerate(ORDER)}
for room,x in positions.items():
    sa_paths.append(([(x+18,120),(x+18,242)],6))
    if ROOMS[room]['rn']:ra_paths.append(([(x+49,302),(x+49,420)],6))
# The last supply main ends at its last branch: no dead extension.
sa_paths[0]=([(115,120),(positions['01']+18,120)],8)
ra_paths[0]=([(115,420),(max(positions[r]+49 for r in RA_ORDER),420)],8)
SA_POLY=network(s,sa_paths,'SA');RA_POLY=network(s,ra_paths,'RA')
assert SA_POLY.intersection(RA_POLY).area<.001
s.line([(106,120),(106,373)],'SA',.5);s.line([(106,120),(115,120)],'SA',.5)
s.line([(115,420),(23,420),(23,174)],'RA',.5)
for y in [133,245,357]:
    s.line([(99,y+16),(106,y+16)],'SA',.5);s.line([(23,y+41),(31,y+41)],'RA',.5)
pair_arrow(s,(106,260),(106,240),'SA');pair_arrow(s,(23,300),(23,280),'RA')
for i,(room,x) in enumerate(positions.items()):
    r=ROOMS[room];s.rect(x,200,68,135,'ARCH',None,.3)
    s.text(x+34,212,'GF-'+room+' / '+SHORT[room],2.65,'INK',True,'center')
    s.damper(x+18,149,'VCD-S'+room,'SA',size=7);s.damper(x+18,178,'MFD-S'+room+'*','SA',motor=True,size=7,fire=True)
    s.text(x+18,193,'RS-SA'+room,2.4,'SA',False,'center',mask=True)
    s.terminal(x+18,242,'SD GROUP','SD',10,'SA');s.text(x+34,263,f"{r['n']} SD / {r['sa']:.2f} L/s",2.5,'SA',False,'center')
    if r['rn']:
        s.terminal(x+49,302,'RG GROUP','RG',10,'RA');s.text(x+34,283,f"{r['rn']} RG / {r['ra']:.2f} L/s",2.5,'RA',False,'center')
        s.text(x+49,341,'RS-RA'+room,2.4,'RA',False,'center',mask=True)
        s.damper(x+49,360,'RCD-R'+room,'RA',motor=True,size=7);s.damper(x+49,393,'MFD-R'+room+'*','RA',motor=True,size=7,fire=True)
        pair_arrow(s,(x+49,316),(x+49,334),'RA')
    else:s.para(x+8,283,'DEDICATED EXTRACT ONLY / MDR-003',52,2.6,'EA',True)
    pair_arrow(s,(x+18,219),(x+18,230),'SA')
    main=next(m for m in MAINS if m['air_type']=='SA' and m['room_takeoff']==room)
    s.text(x+18,104,main['tag'],2.5,'SA',True,'center')
    s.text(x+18,92,f"{main['flow_before_takeoff_l_s']:.1f} L/s",2.45,'SA',False,'center')
    if i<8:pair_arrow(s,(x+36,120),(x+58,120),'SA')
for room in RA_ORDER:
    x=positions[room];m=next(m for m in MAINS if m['air_type']=='RA' and m['room_takeoff']==room)
    s.text(x+49,441,m['tag']+f" / {m['flow_before_takeoff_l_s']:.1f} L/s",2.45,'RA',False,'center')
    pair_arrow(s,(x+38,420),(x+15,420),'RA')
heading(s,20,471,'HOW THE METHOD WORKS',800)
note(s,20,486,'DX cooling takes place inside the PAUs. Refrigerant stays within the selected factory DX system; supply and return mains carry AIR. Common SUPPLY and RETURN mains stay on the ROOF. Nine supply drops and six return risers connect room branches to the shared headers; every PAU connects to the same header bank. Terminal groups and all instruments are expanded on the following sheets.',800)
note(s,20,509,'Supply 3697.80 L/s; assumed return 3006.68 L/s; minimum mass-balance makeup 691.12 L/s. DX cooling capacity / outside air / actual pressures remain unselected. Rooms with dedicated extract have no common return branch.',800,2.6)
footer(s,'NTS functional air diagram: room order and sequential flow steps are illustrative; actual grouped roof takeoffs / coordinates on MDR-012. MFD* conditional on approved fire boundaries; RCD pressure modulation is a proposal. All room flows retained. Main sizes: MDR-011.');SHEETS.append(s)

# MDR-002: every terminal tag, separate supply and return headers within each room.
s=frame(2,'ALL ROOM DIFFUSERS, RETURN GRILLES AND EXTRACT PICKUPS')
for i,room in enumerate(ORDER):
    x=20+(i%3)*271;y=45+(i//3)*160;r=ROOMS[room]
    s.rect(x,y,258,149,'DIM',None,.3);s.text(x+7,y+11,'GF-'+room+' '+SHORT[room],3,'INK',True)
    for air,basey in [('SA',y+54),('RA',y+105),('EA',y+105)]:
        terms=[t for t in TERMS if t['room']==room and t['air_type']==air]
        if not terms:continue
        count=len(terms);xs=[x+22+j*(214/max(1,count-1)) for j in range(count)] if count>1 else [x+125]
        hy=basey-17 if air=='SA' else basey+29
        paths=[([(x+8,hy),(xs[-1],hy)],4)]
        for xx,t in zip(xs,terms):paths.append(([(xx,hy),(xx,basey)],3.5))
        network(s,paths,air)
        pair_arrow(s,(x+9,hy),(x+18,hy),air) if air=='SA' else pair_arrow(s,(x+20,hy),(x+9,hy),air)
        for xx,t in zip(xs,terms):
            s.terminal(xx,basey,'','SD' if air=='SA' else 'RG',8,air)
            label_x=xx if air=='SA' else xx-6
            align='center' if air=='SA' else 'right'
            s.text(label_x,basey+12,t['tag'],2.25,air,False,align)
            q=numeric(t['flow_l_s']);s.text(label_x,basey+22,(f'{q:.2f} L/s' if q is not None else 'FLOW TBC'),2.2,air,False,align)
        s.text(x+8,hy-5 if air=='SA' else hy+10,'FROM SA MAIN' if air=='SA' else 'TO RA MAIN' if air=='RA' else 'TO DEDICATED EF',2.25,air)
footer(s,'All 53 terminal tags / duties are retained. Glyph represents the registered terminal face; this sheet is a functional room header diagram, not ceiling setting-out. Terminal neck / face sizes and coordinates retained in the workbook.');SHEETS.append(s)

# MDR-003: contaminated / process extracts remain outside recirculated return.
s=frame(3,'DEDICATED EXTRACTS AND RETURN AIR SEPARATION')
for i,room in enumerate(['03','09','01']):
    x=20+i*271;r=ROOMS[room];s.rect(x,49,258,397,'DIM',None,.3)
    s.text(x+8,62,'GF-'+room+' '+SHORT[room],3.2,'INK',True)
    s.rect(x+16,82,221,53,'ARCH',None,.35);s.text(x+126,103,'NO CONNECTION TO COMMON RETURN',2.8,'EA',True,'center')
    s.text(x+126,121,'EG pickup -> dedicated extract duct',2.8,'EA',False,'center')
    qty=1 if room=='01' else 2
    fan_levels=[173+j*112 for j in range(qty)]
    outlet_y=342 if qty==2 else 249
    paths=[([(x+126,135),(x+126,155),(x+15,155),(x+15,fan_levels[-1])],9),
           ([(x+240,fan_levels[0]),(x+240,outlet_y)],9)]
    paths += [([(x+15,yy),(x+240,yy)],9) for yy in fan_levels]
    network(s,paths,'EA')
    pair_arrow(s,(x+126,137),(x+126,150),'EA')
    pair_arrow(s,(x+240,outlet_y-23),(x+240,outlet_y-7),'EA')
    s.text(x+230,outlet_y+13,'TO EXHAUST DISCHARGE',2.55,'EA',True,'right')
    for j in range(qty):
        yy=173+j*112;tag=f'EF-{room}-{j+1:02}'
        s.flex(x+48,yy,'EA');s.fan(x+107,yy,tag,'EA',10);pair_arrow(s,(x+128,yy),(x+153,yy),'EA')
        s.damper(x+177,yy,'MD-E'+room+'-'+f'{j+1:02}','EA',motor=True,size=9)
        sensor(s,x+83,yy-29,'DPT-'+tag,3,'DP');sensor(s,x+215,yy-29,'VS-'+tag,3,'VS')
        s.line([(x+83,yy-24),(x+74,yy-24),(x+74,yy),(x+98,yy)],'INST',.25)
        s.line([(x+87,yy-29),(x+135,yy-29),(x+135,yy),(x+117,yy)],'INST',.25)
        s.line([(x+211,yy-29),(x+196,yy-15),(x+113,yy-5)],'INST',.25)
        q='TBC' if room=='01' else f"{float(r['ea']):.2f} L/s"
        s.text(x+112,yy+32,'EACH DUTY FAN: '+q,2.65,'EA',False,'center')
    note(s,x+9,375,'One duty / one standby; motorized isolation and run/fault changeover. Discharge location, dirty-air separation and losses require coordination.' if qty==2 else 'Clean-agent room normal extract duty / fan quantity / pickup elevation / agent suitability remain TBC. No automatic fire extraction is added.',239,2.8)
heading(s,20,472,'AIR BALANCE / FIRE REQUIREMENTS',800)
note(s,20,487,'Toilet SA 46.20 minus EA 64.89 = -18.69 L/s; kitchen SA 147.60 minus EA 147.60 = zero. These do not establish the previous +25 Pa targets. Agree transfer / hygiene / pressure strategy using measured leakage and project requirements.',795)
footer(s,'Confirmed fire stops every HVAC fan and heater; fire reset is manual. No automatic smoke extraction or post-fire purge is implied. Dedicated extract fan outlet / discharge duty and CAG selection remain open.');SHEETS.append(s)

# MDR-004: the unit paths merge into common headers; all 25 air-side instruments.
s=frame(4,'DX PAU BANK, COMMON HEADERS AND AIR-SIDE INSTRUMENTS')
s.line([(28,65),(28,425)],'RA',.6);s.line([(790,65),(790,425)],'SA',.6)
s.text(27,53,'COMMON RA',2.8,'RA');s.text(790,53,'COMMON SA',2.8,'SA',False,'right')
for i,c in enumerate('ABC'):
    y=62+i*122;cy=y+59;s.rect(44,y,730,108,'DIM',None,.25)
    s.text(51,y+11,'PAU-'+c+' | 1848.90 L/s SA | 2 DUTY / 1 STANDBY BANK',2.85,'INK',True)
    s.line([(28,cy),(790,cy)],'SA',.42);s.line([(28,cy),(94,cy)],'RA',.45)
    s.damper(93,cy,'MD-RA-'+c,'RA',motor=True,size=9)
    s.line([(130,y+18),(130,cy)],'OA',.45);s.damper(130,y+36,'MD-OA-'+c,'OA',motor=True,size=7)
    s.text(158,y+24,'OA branch: mass-balance >=345.56 L/s per duty unit',2.5,'OA')
    for j,fx in [(1,232),(2,346)]:
        s.filter(fx-5,cy-13,10,26,'F'+str(j)+'-'+c)
        sensor(s,fx,cy-28,f'PDT-F{j}-{c}',4,'DP');sensor(s,fx+24,cy+30,f'PDI-F{j}-{c}',4,'DP')
        for xx,sgn in [(fx-9,-1),(fx+9,1)]:
            s.line([(xx,cy),(xx,cy-28),(fx+sgn*4.5,cy-28)],'INST',.2)
            s.line([(xx,cy+4),(xx,cy+30),(fx+24+sgn*4.5,cy+30)],'INST',.2)
    s.rect(430,cy-14,32,28,'SA','#FFFFFF',.4)
    for dx in [434,443,452]:s.line([(dx,cy-10),(dx+5,cy+10)],'RL',.4)
    s.text(446,cy+34,'DX COIL / MDR-005',2.4,'RL',True,'center')
    s.fan(566,cy,'SF-'+c,'SA',10);s.flex(636,cy,'SA');s.damper(716,cy,'MD-SA-'+c,'SA',motor=True,size=9)
    sensor(s,545,cy-28,'DPT-SF-'+c,4,'DP');sensor(s,620,cy-28,'VS-PAU-'+c,4,'VS')
    s.line([(541,cy-28),(520,cy-28),(520,cy),(555,cy)],'INST',.2)
    s.line([(549,cy-28),(589,cy-28),(589,cy),(577,cy)],'INST',.2)
    s.line([(616,cy-28),(588,cy-8)],'INST',.2)
    pair_arrow(s,(482,cy),(507,cy),'SA')
pair_arrow(s,(28,405),(28,377),'RA');pair_arrow(s,(790,376),(790,408),'SA')
# OA filter / two instruments.
s.duct([(25,484),(190,484)],8,'OA');s.filter(96,473,9,22,'F-OA-01','OA')
sensor(s,86,453,'PDT-OA-01',4,'DP');sensor(s,142,511,'PDI-OA-01',4,'DP')
for px,sgn in [(90,-1),(111,1)]:
    s.line([(px,484),(px,453),(86+sgn*4.5,453)],'INST',.2)
    s.line([(px,487),(px,511),(142+sgn*4.5,511)],'INST',.2)
s.line([(228,464),(806,464)],'SA',.5);sensor(s,349,464,'FIT-SA-01',4,'AF');sensor(s,670,464,'TT-SA-01',4,'TT')
s.line([(806,514),(228,514)],'RA',.5)
for x,t,l in [(349,'FIT-RA-01','AF'),(537,'TT-RA-01','TT'),(708,'RHT-RA-01','RH')]:sensor(s,x,514,t,4,l)
footer(s,'Standby PAU supply / return / OA paths close and prove isolated; duty paths open and prove before start. Common airflow stations require OEM straight runs. Unit nozzle routing, fan ESP, filters, cooling capacity and outdoor air remain unselected.');SHEETS.append(s)

# MDR-005: the DX refrigeration cycle is local to each PAU, with separate new proposals.
s=frame(5,'DX REFRIGERANT CYCLE AND PROPOSED OEM PROTECTION')
for i,c in enumerate('ABC'):
    x=20+i*271;s.rect(x,48,258,302,'DIM',None,.3);s.text(x+8,62,'PAU-'+c+' - FACTORY DX CIRCUIT',3.1,'INK',True)
    cx,cond=x+67,x+200
    s.line([(cx,109),(cond,109),(cond,238),(cx,238),(cx,109)],'RL',.6)
    s.rect(cx-19,98,38,22,'INK','#FFFFFF',.35);s.text(cx,112,'COMPRESSOR',2.5,'INK',True,'center')
    s.rect(cond-21,98,42,23,'INK','#FFFFFF',.35);s.text(cond,113,'CONDENSER',2.5,'INK',True,'center')
    s.arrow((cond+22,110),(x+245,110),'EA',.25,1.8);s.text(x+206,91,'HEAT TO OUTDOORS',2.3,'EA',False,'center')
    s.rect(cond-9,172,18,18,'INK','#FFFFFF',.3);s.text(cond,184,'FD',2.5,'INK',True,'center');s.text(cond-13,197,'FILTER / DRIER',2.3)
    s.circle(cond,218,5,'RL',.3,'#FFFFFF');s.text(cond,219,'EEV',2.3,'RL',True,'center')
    s.rect(cx-26,226,52,24,'SA','#FFFFFF',.35);s.text(cx,240,'EVAPORATOR',2.6,'SA',True,'center')
    pair_arrow(s,(cx+25,109),(cx+46,109),'RL');pair_arrow(s,(cond,144),(cond,160),'RL');pair_arrow(s,(cond-30,238),(cond-56,238),'RL');pair_arrow(s,(cx,153),(cx,135),'RL')
    sensor(s,x+130,81,'PSH-DX-'+c,5,'HP',True);s.line([(x+130,85.5),(x+130,109)],'INST',.2)
    sensor(s,x+35,161,'PSL-DX-'+c,5,'LP',True);s.line([(x+39.5,161),(cx,161)],'INST',.2)
    sensor(s,x+229,161,'PT-LIQ-DX-'+c,5,'PT',True);s.line([(x+224.5,161),(cond,161)],'INST',.2)
    sensor(s,x+114,204,'PT-SUC-DX-'+c,5,'PT',True);s.line([(x+109.5,204),(cx,204)],'INST',.2)
    sensor(s,x+38,280,'TSLL-DX-'+c,5,'LT',True);s.line([(x+38,275.5),(cx-10,250)],'INST',.2)
    sensor(s,x+134,280,'LSH-DX-'+c,5,'LS',True);s.line([(x+134,275.5),(cx+20,250)],'INST',.2)
    sensor(s,x+35,81,'OL-CMP-DX-'+c,5,'OL',True);s.line([(x+35,85.5),(cx-19,100)],'INST',.2)
    s.text(x+128,328,'REFRIGERANT / CHARGE / PIPE SIZES: OEM',2.7,'PENDING',False,'center')
heading(s,20,376,'OPTIONAL COMMON-MAIN STATIC CONTROL / NEW PROPOSAL',800)
s.duct([(35,449),(348,449)],10,'SA');sensor(s,181,421,'DPT-SA-M01',5,'SP',True);s.line([(181,425.5),(181,449)],'INST',.25)
s.line([(185.5,421),(395,421),(395,445),(442,445)],'CTRL',.3,True)
s.rect(442,417,150,57,'INST',None,.35);s.text(517,438,'HPCP STATIC CONTROL',2.8,'INST',True,'center');s.text(517,456,'OEM FAN / VFD SUITABILITY TBC',2.5,'INST',False,'center')
s.line([(592,445),(656,445)],'CTRL',.3,True);s.rect(656,417,158,57,'SA',None,.35);s.text(735,445,'DUTY PAU FAN CONTROL',2.8,'SA',True,'center')
note(s,20,490,'22 new DX protection / monitoring candidates are separate from the 85 retained air-side instruments. Factory compressor / condenser / evaporator control governs. The shown circuit explains DX operation; receiver, accumulator, refrigerant, valve, oil return, relief and electrical arrangements follow the selected OEM circuit.',797)
footer(s,'No field refrigerant network to the rooms is proposed. Cooling capacity, refrigerant pressure-temperature ranges and compressor quantity/staging are TBC. Do not modify factory safety wiring from this diagram.');SHEETS.append(s)

# MDR-006: every one of the 27 room TT/RHT/DPT requirements graphically referenced.
s=frame(6,'ROOM TEMPERATURE, HUMIDITY AND PRESSURE INSTRUMENTS')
for i,room in enumerate(ORDER):
    x=20+i%3*271;y=48+i//3*158;s.rect(x,y,258,143,'DIM',None,.3)
    s.text(x+8,y+13,'GF-'+room+' '+SHORT[room],3,'INK',True)
    for xx,pre,label in [(x+45,'TT','TT'),(x+123,'RHT','RH'),(x+211,'DPT','DP')]:sensor(s,xx,y+49,pre+'-GF-'+room,6,label)
    s.line([(x+206.5,y+49),(x+173,y+49),(x+173,y+76)],'INST',.25)
    s.text(x+167,y+89,'(+) ROOM STATIC',2.45,'INST')
    s.line([(x+215.5,y+49),(x+246,y+49),(x+246,y+108)],'INST',.25)
    s.text(x+246,y+120,'(-) OUTDOOR',2.4,'INST',False,'right')
    s.text(x+11,y+88,'TT / RH at +1.50 m AFFL proposed',2.55)
    s.text(x+11,y+103,'3 AI to HPCP-01; room flow balanced',2.55,'INST')
    s.text(x+11,y+119,'Pressure target '+str(int(ROOMS[room]['pressure']))+' Pa / not verified',2.55,'PENDING')
    s.text(x+11,y+134,'Avoid supply jets / solar heat / heater faces',2.45,'DIM')
footer(s,'Room pressure uses a sheltered outdoor reference, equal datums and static pickups. RCD control is a proposed pressure trim; TT/RHT monitor the room and do not create independent DX zones. Final pressure, comfort control, ranges, setpoints and tubing are pending.');SHEETS.append(s)

# MDR-007: all six AFS / high limits, including common heater, remain explicit.
s=frame(7,'HEATER PERMISSIVES, ACCESS AND HARDWIRED SAFETY')
for i,suffix in enumerate(['04','05','06','07','03','COM-01']):
    y=54+i*72;tag='EH-'+suffix;s.text(22,y+8,tag+' | RETAINED HEATING PROVISION',2.8,'INK',True)
    s.duct([(31,y+35),(493,y+35)],11,'SA')
    if suffix!='COM-01':s.damper(89,y+35,'VCD-S'+suffix,'SA',size=8)
    s.heater(283,y+24,tag,'SA',39,23)
    sensor(s,180,y+12,'AFS-EH-'+suffix,7,'AF');s.line([(180,y+16.5),(180,y+35)],'INST',.25)
    sensor(s,382,y+12,'TSHH-EH-'+suffix,7,'HL');s.line([(382,y+16.5),(315,y+26)],'INST',.25)
    s.access(257,y+42,'AD-H'+suffix if suffix!='COM-01' else 'AD-HCOM-01',size=6)
    s.text(450,y+57,'SAFE HEAT ENABLE',2.5,'INST',False,'right')
heading(s,540,52,'OEM SAFETY CHAIN / INTERFACE',278)
for i,t in enumerate(['FIRE TRIP / E-STOP HEALTHY','DUTY FAN PROOF','AFS AIRFLOW PROVED','TSHH HIGH LIMIT HEALTHY','OEM CONTACTOR / HEATING ENABLE']):
    y=79+i*49;s.rect(550,y,258,31,'INST',None,.35);s.text(679,y+18,t,2.8,'INST',True,'center')
    if i<4:s.arrow((679,y+31),(679,y+46),'INST',.25,2)
note(s,540,347,'Loss of airflow or high-limit trip inhibits heat through the independent OEM circuit. HPCP heating demand cannot bypass either safety. Contacts and final electrical terminals follow the selected heater manufacturer.',275)
note(s,540,420,'DX provides cooling. The retained room/common electric heaters and conditional HUM-01 serve distinct heating/humidity duties. Omitting them from a cooling-only operation requires confirmation of the project thermal/humidity requirements; no automatic deletion is made.',275)
footer(s,'Six AFS and six independent TSHH provisions are retained. Heater dimensions, power, thermal clearances, airflow proof technology and contact reset are OEM selections. Dampers/access symbols are functional; MFD quantity/rating follows the fire strategy.');SHEETS.append(s)

# MDR-008: roof crossing and retained room-drop ceiling envelopes.
s=frame(8,'ROOF CROSSING, ROOM DROPS AND CEILING ENVELOPES')
heading(s,20,47,'A  ROOF BRANCH CROSSING / RCX-01 | VERTICAL 1:20',390)
heading(s,434,47,'B  ROOM DUCT + FLANGES | SECTION 1:10',386)
y=253
# Actual RCX-01: 700x450 SA split branch under 700x900 RA main.
for x,zlo,zhi,w,air in [(105,4800,5250,35,'SA'),(68,5550,6450,130,'RA')]:
    top=y-(zhi-3300)/20;hh=(zhi-zlo)/20
    s.rect(x,top,w,hh,air,None,.5);s.rect(x-3.75,top-3.75,w+7.5,hh+7.5,'DIM',None,.3)
s.line([(25,y-(4200-3300)/20),(405,y-(4200-3300)/20)],'ARCH',.45)
s.text(263,105,'RA BOD +5.550 m',2.8,'RA');s.text(263,159,'SA BOD +4.800 m',2.8,'SA')
s.text(211,138,'150 mm BETWEEN INSULATION',2.7);s.text(263,202,'ROOF +4.200 m ASSUMED',2.7,'DIM')
s.text(31,242,'75 mm OUTDOOR INSULATION EACH FACE ASSUMED',2.65)
# Above-ceiling room ducts remain single layer; common roof mains do not enter this void.
cy=240;s.line([(441,cy),(817,cy)],'ARCH',.4);s.text(442,cy+11,'CEILING +3.300 m',2.7)
s.rect(441,cy-72,376,2,'ARCH',None,.4);s.text(654,cy-76,'SLAB U/S +4.000 m',2.7)
s.rect(479,cy-57.5,119,55,'INK',None,.35);s.rect(481.5,cy-55,114,50,'DIM',None,.3);s.rect(486.5,cy-50,104,40,'SA',None,.5)
s.text(622,cy-49,'BARE DEPTH <=400 mm',2.7,'SA');s.text(622,cy-33,'50 mm INDOOR INSULATION*',2.7)
s.text(622,cy-17,'25 mm FLANGE EACH SIDE*',2.7);s.text(622,cy-3,'CEILING GAP 25 / SLAB 125 mm',2.6)
heading(s,20,291,'15 ROOM DROPS: 9 SUPPLY + 6 RETURN / PROPOSED LEVELS',552)
risers=read('Roof_risers.csv')
table(s,20,304,552,['Riser','Room','Clear W x H','Indoor BOD','Roof BOD','Rise m*'],[[r['tag'],r['room'],r['width_mm']+'x'+r['height_mm'],'+3.400',f"{float(r['roof_bod_mm_affl'])/1000:.3f}",r['vertical_length_m']] for r in risers],[1.4,.5,1.2,1,1,1],2.65,13)
heading(s,593,291,'CONNECTION / INSTALLATION HOLDS',226)
note(s,593,307,'Roof SA main 600x1300 and RA 700x900 stay outdoors. Only local room ducts descend into the 700 mm ceiling void. Roof-header plan separation and registered crossing levels are shown on MDR-012.',224,2.65)
note(s,593,379,'Each RS tag matches its room branch. Two fitted elbows / roof transition, curb / weather seal / firestop and branch access require shop development. Roof +4.200 m assumes a 200 mm slab above the confirmed +4.000 m underside.',224,2.65)
note(s,593,467,'* Proposed envelope / rise only. Actual roof structure, beam downstands, insulation specification, actuators, heaters, supports and plenums require OEM/site coordination.',224,2.6)
footer(s,'User selected roof common mains with room drops. A shows an actual registered roof branch crossing, not stacked indoor common mains; B uses the retained deepest local-room envelope. Insulation and flange dimensions are assumptions.');SHEETS.append(s)

# MDR-009: retain separate basement, power, all 11 auxiliary instruments and controls.
s=frame(9,'BASEMENT, HPCP, ENERGY METERING AND CONTROL SEQUENCE')
for i,tag in enumerate(['SF-B1-01','SF-B1-02','EF-B1-01','EF-B1-02']):
    y=79+i*100;air='SA' if tag.startswith('SF') else 'EA'
    s.duct([(35,y),(500,y)],11,air);s.fan(209,y,tag,air,10);s.damper(370,y,'MD-'+tag,air,motor=True,size=9)
    sensor(s,145,y-27,'DPT-'+tag,9,'DP');sensor(s,292,y-27,'VS-'+tag,9,'VS')
    s.line([(140.5,y-27),(115,y-27),(115,y),(197,y)],'INST',.23)
    s.line([(149.5,y-27),(242,y-27),(242,y),(220,y)],'INST',.23)
    s.line([(287.5,y-27),(256,y-11),(215,y-5)],'INST',.23)
    s.text(35,y+34,'2 x 100% duty / standby; basement flow / heat / volume TBC',2.7)
heading(s,540,45,'BASEMENT / POWER INSTRUMENTS',279)
sensor(s,595,91,'TT-B1-01',9,'TT');sensor(s,747,91,'DPT-B1-01',9,'DP')
s.text(559,116,'Pressure (+) basement / (-) sheltered outdoor',2.55,'INST')
s.rect(551,145,257,98,'INST',None,.35);sensor(s,679,179,'EM-HPCP-01',9,'EM')
s.text(680,211,'V / I / kW / kWh / PF / Hz / VFD THD',2.7,'INST',False,'center')
s.text(679,229,'Feeder >7.5 kW; ratings / CT / protocol TBC',2.6,'INST',False,'center')
heading(s,540,266,'HPCP-01 / HMI / BMS SEQUENCE',279)
yy=282
for t in ['Start: healthy fire/ESD and OEM safety; select the duty PAU pair; open/prove RA, OA and SA dampers before enabling selected fans/compressors.',
          'Prove fan operation; verify common airflow; enable DX through the OEM controller. Heating requires separate hardwired AFS/TSHH healthy permissives.',
          'Rotate duty pair; failed unit stops / isolates, standby opens/proves then starts. Avoid reverse airflow through an idle unit.',
          'Confirmed fire stops all fans/heaters and latches ESD; manual reset after clearance. Unit compressor response and damper fire positions per approved OEM/fire design.']:
    yy+=note(s,540,yy,t,275,2.65)
counts={t:sum(r['type']==t for r in IO) for t in ['AI','AO','DI','DO','COMM']}
table(s,20,471,490,['Retained I/O','AI','AO','DI','DO','COMM'],[['275 baseline points',*[counts[t] for t in counts]]],[2,1,1,1,1,1],2.7,16)
footer(s,'Basement remains an independent ventilation system; its flows/ESP/heat removal are not included in the 3697.80 L/s room supply. New DX proposals are separate from the unchanged 275-row I/O baseline; final channel allocation and protocols pending.');SHEETS.append(s)

# MDR-010: cross-reference every retained instrument to its actual graphical sheet.
assert set(COVERED)=={r['tag'] for r in INSTRUMENTS},('Uncovered retained instrument',set(r['tag'] for r in INSTRUMENTS)-set(COVERED))
assert set(DX_COVERED)=={r['tag'] for r in DX}
index=[dict(tag=r['tag'],parameter=r['parameter'],location=r['location'],io=r['io'],
            graphical_sheet=' / '.join(f'MDR-{n:03}' for n in sorted(COVERED[r['tag']])),status='Retained requirement / settings and OEM wiring pending') for r in INSTRUMENTS]
write('Instrument_drawing_index.csv',index)
s=frame(10,'COMPLETE INSTRUMENT INDEX AND DRAWING REFERENCES')
for offset,x in [(0,20),(43,427)]:
    table(s,x,49,394,['Retained tag','Parameter','Graphic ref / I/O'],[[r['tag'],r['parameter'],r['graphical_sheet']+' / '+r['io']] for r in index[offset:offset+43]],[1.8,2.6,1.5],2.35,10.4)
footer(s,'85 retained instruments are graphically referenced on MDR-003 through MDR-009. 22 new DX candidates are on MDR-005 and in Additional_DX_proposals. Original instrument requirements and I/O are unchanged; final OEM selection / safety settings remain open.');SHEETS.append(s)

# MDR-011: 15 functional roof-header flow steps, OA balance and fitting development.
s=frame(11,'COMMON HEADER FLOW ACCOUNTING AND ENGINEERING BASIS')
table(s,20,48,800,['Main','Takeoff GF','Before L/s','Branch L/s','After L/s','Clear W x H','m/s','Straight Pa/m'],[[r['tag'],r['room_takeoff'],f"{r['flow_before_takeoff_l_s']:.2f}",f"{float(r['takeoff_l_s']):.2f}",f"{r['flow_after_takeoff_l_s']:.2f}",str(r['clear_width_mm'])+'x'+str(r['clear_depth_mm']),f"{r['velocity_m_s']:.3f}",f"{r['straight_friction_pa_m']:.3f}"] for r in MAINS],[1.2,.9,1.2,1.2,1.2,1.3,.8,1.2],2.75,16)
heading(s,20,329,'CALCULATION / FABRICATION BASIS',389)
yy=343
for t in ['Retained roof spines: SA600x1300, RA700x900. This sequential room flow accounting explains the method; actual grouped roof route / branch dimensions and duties are authoritative in Roof_routes_Rev12 and MDR-012. Colebrook/Darcy-Weisbach, rho1.2, mu1.81e-5, galvanized roughness0.09 mm. These are preliminary design choices, not selected acoustic performance.',
          'At each room SA takeoff, remaining flow decreases. On RA the sequence is listed away from the PAU; actual airflow collects toward the PAU. The scaled roof uses two grouped SA branch columns, so these sequential steps are not a fabrication main schedule. No toilet/kitchen/CAG return enters this header.',
          'Develop each reducer with L >= max(300, max(|delta W|,|delta H|)/(2 tan15deg)). Flared three-leg boots and open throats; selected turning vanes/radii and device face-to-face lengths govern shop drawings.']:
    yy+=note(s,20,yy,t,387,2.7)
heading(s,434,329,'THERMAL / PRESSURE / SELECTION HOLD',386)
yy=343
for t in [f'SA {TOTAL_SA:.2f} minus RA {TOTAL_RA:.2f} requires makeup >= {MAKEUP:.2f} L/s. OA reference420.70 is short by {MAKEUP-420.7:.2f} L/s. This is a mass-balance lower bound, not an approved outdoor-air duty.',
          'No complete fan ESP is calculated for this method: surveyed common-main/riser lengths, nozzle losses, branch layout, fitted K, terminals, coils/dirty filters, intake and exhaust discharge are unselected. The Rev12 allowance budget applies only to its measured proposed room-drop routes, excludes PAU bank/nozzle/OA details and does not select a fan for this method package.',
          'DX coil total/sensible/latent capacity and refrigerant must be selected from revised room/OA loads and design outdoor conditions. One common supply condition alone does not prove independent room comfort; verify loads, reheat/humidity requirements and zoning.']:
    yy+=note(s,434,yy,t,383,2.7)
footer(s,'Review index: 001 whole method; 002 all terminals; 003 extracts; 004 PAU/air instruments; 005 DX circuit; 006 room instruments; 007 heater safety; 008 sections/risers; 009 controls/basement; 010 all 85 instruments; 011 functional flow steps; 012 scaled roof plan; 013 indoor coordination; 014 room instrument locations.');SHEETS.append(s)

assert len(SHEETS)==11
def export_pdf(name,mono=False):
    old=COL.copy()
    if mono:COL.update({k:'#AAAAAA' if k=='ARCH' else '#000000' for k in COL})
    sheets=SHEETS
    if mono:
        sheets=[]
        for original in SHEETS:
            sheet=Scene()
            for operation in original.ops:
                op=list(operation)
                fill_slot=3 if op[0]=='poly' else 6 if op[0]=='circle' else None
                if fill_slot is not None and op[fill_slot]:
                    fill=op[fill_slot];rgb=[int(fill[j:j+2],16) for j in [1,3,5]]
                    if fill!='#FFFFFF':op[fill_slot]='#EEEEEE' if min(rgb)>200 else '#000000'
                sheet.ops.append(tuple(op))
            sheets.append(sheet)
    c=canvas.Canvas(str(OUT/name),pagesize=(841*MM,594*MM),pageCompression=1)
    c.setTitle(TITLE);c.setAuthor('Administration HVAC engineering coordination');c.setSubject('DX common supply / common return method; provisional engineering inputs')
    for sheet in sheets:pdf_render_scene(c,sheet);c.showPage()
    c.save();COL.update(old)
export_pdf(TITLE+'.pdf');export_pdf(TITLE+' - Monochrome.pdf',True)
export_schematic(SHEETS,OUT/(TITLE+'.dxf'))
# Keep the selected Rev12 physical plans as clearly identified references.
reference_pages=[22,0,6]
for filename in [TITLE+'.pdf',TITLE+' - Monochrome.pdf']:
    with fitz.open(OUT/filename) as pdf, fitz.open(BASE/'inputs'/('Rev12_reference_mono.pdf' if 'Monochrome' in filename else 'Rev12_reference.pdf')) as ref:
        for n,page in enumerate(reference_pages,12):
            pdf.insert_pdf(ref,from_page=page,to_page=page)
            pdf[-1].insert_text((22*MM,36*MM),f'MDR-{n:03} REFERENCE PLAN / RETAINED REV12 D-{page+1:03} / MAIN DUCT AND RETURN',fontsize=2.5*MM,color=(.25,.25,.25))
        pdf.save(OUT/(filename+'.tmp'))
    (OUT/(filename+'.tmp')).replace(OUT/filename)
# Native CAD references use their original model blocks, translated away from the NTS diagrams.
from ezdxf.xref import Loader,ConflictPolicy
from ezdxf.math import Matrix44
target=ezdxf.readfile(OUT/(TITLE+'.dxf'));source=ezdxf.readfile(BASE/'inputs/Rev12_reference.dxf')
# Discard only stale frozen-layer names, keeping valid reference visibility settings.
for original_layout in source.layouts:
    for vp in original_layout.query('VIEWPORT'):
        vp.frozen_layers=[name for name in vp.frozen_layers if name in source.layers]
# Presentation corrections stay in memory; original Rev12 input is unchanged.
for original_block in source.blocks:
    for text in original_block.query('TEXT'):text.dxf.height*=.718
for name in ['D023_A1','D001_A1','D007_A1']:
    for hatch in source.layouts.get(name).query('HATCH[layer=="M-ANNOTATION"]'):
        hatch.dxf.true_color=0xF0F3F5
existing={e.dxf.handle for e in target.modelspace()}
# Modern loader preserves WIPEOUTs and separate source dash/style resources.
loader=Loader(source,target,conflict_policy=ConflictPolicy.NUM_PREFIX);loader.load_modelspace()
ref_layouts=[]
for n,name in enumerate(['D023_A1','D001_A1','D007_A1'],12):
    loader.load_paperspace_layout(source.layouts.get(name));ref_layouts.append((name,f'MDR{n:02}_A1'))
loader.execute()
for e in target.modelspace():
    if e.dxf.handle not in existing:e.transform(Matrix44.translate(900000,0,0))
for original_name,newname in ref_layouts:
    layout=target.layouts.get(original_name)
    target.layouts.rename(layout.name,newname)
    for vp in layout.query('VIEWPORT'):
        if vp.dxf.id>1:
            center=vp.dxf.view_center_point;vp.dxf.view_center_point=(center.x+900000,center.y)
    layout.add_text('REFERENCE PLAN / RETAINED REV12 / MAIN DUCT AND RETURN',dxfattribs={'height':2.5*.718,'layer':'M-INK'}).set_placement((22,558))
target.saveas(OUT/(TITLE+'.dxf'))

# Workbook keeps all source duty data and distinguishes new DX candidates.
wb=Workbook();ws=wb.active;ws.title='Read_me'
for row in [('Drawing package',TITLE+' / 11 method sheets + 3 scaled reference plans / 85 retained instruments + 22 new DX candidates'),
            ('Architecture','DX PAU bank -> roof common SA main -> nine room drops; six room returns -> roof common RA main -> bank. Dedicated extracts separate.'),
            ('Main location','User selected ROOF COMMON MAINS WITH ROOM DROPS. Scaled roof plan is a floor-footprint proposal; roof survey and OEM approval pending.'),
            ('Ceiling','User slab U/S4000, ceiling3300 mm. Indoor room depth <=400 mm, BOD3400. Roof SA BOD4800, RA5550 proposed; 50 indoor/75 outdoor insulation assumptions.'),
            ('Scope','Method drawings / sizing only; actual building routing, supports, structure, fan ESP, DX capacity and OEM selections pending.'),
            ('AutoCAD','Open Main duct and return.dxf; MDR01_A1 through MDR14_A1; Save As DWG. Native AutoCAD plotting has not run.'),
            ('Input editing','Original inputs retained; edit build.py for the method assumptions. Generated workbook edits are not imported.')]:ws.append(row)
mapping=[('Instruments.csv','Retained_instruments'),('IO_points.csv','Retained_IO'),('Terminals.csv','All_terminals'),('Components.csv','Retained_components'),('Instrument_installation_Rev12.csv','Mounting_baseline'),('Roof_routes_Rev12.csv','Roof_routes_Rev12'),('Roof_risers.csv','Roof_risers')]
for name,title in mapping:
    ws=wb.create_sheet(title)
    with (BASE/'inputs'/name).open() as f:
        for row in csv.reader(f):ws.append(row)
for name,title in [('Common_main_sizing.csv','Common_main_sizing'),('Additional_DX_proposals.csv','Additional_DX'),('Instrument_drawing_index.csv','Instrument_index')]:
    ws=wb.create_sheet(title)
    with (OUT/name).open() as f:
        for row in csv.reader(f):ws.append(row)
for ws in wb:
    ws.freeze_panes='A2';ws.sheet_view.showGridLines=False;ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='123B52');c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='top')
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10)
        ws.row_dimensions[row[0].row].height=60 if ws.title in ['Read_me','Additional_DX','Mounting_baseline'] else 38
    for j,col in enumerate(ws.iter_cols(),1):ws.column_dimensions[get_column_letter(j)].width=min(80,max(18,max(len(str(c.value or '')) for c in col)*.65))
    ws.print_title_rows='1:1';ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
wb.save(OUT/(TITLE+' - Registers.xlsx'))

summary={'name':TITLE,'method_revision':1,'sheets':COUNT,'baseline_instruments':85,'new_DX_proposals':22,
         'terminals':53,'baseline_IO_points':len(IO),'common_roof_route_segments':31,'functional_header_flow_steps':len(MAINS),'room_drops':15,
         'supply_l_s':TOTAL_SA,'return_l_s':TOTAL_RA,'minimum_mass_balance_makeup_l_s':MAKEUP,
         'ceiling_clearance_requirement_mm':3300,'slab_underside_mm':4000,'available_void_mm':700,
         'main_location':'Roof common mains with room drops - user selected','bare_indoor_room_duct_depth_limit_mm':400,'assumed_insulation_each_face_mm':50,
         'assumed_flange_allowance_each_side_mm':25,'component_reserved_depth_mm':550,
         'status':'Detailed functional method and preliminary sizing, not surveyed route or construction approval',
         'holds':['Proposed roof layout / structure / penetrations / weatherproofing / support and OEM dimensions','Room drops: beam / actuator / support / access coordination',
                  'DX capacity, refrigerant, OEM circuit and fan selection / complete route ESP','OA/pressure balance and room comfort / humidity requirements','Fire boundaries and final wiring']}
(OUT/'method_basis.json').write_text(json.dumps(summary,indent=2)+'\n')
review=OUT/'review';review.mkdir(exist_ok=True)
with fitz.open(OUT/(TITLE+'.pdf')) as pdf:
    for i,p in enumerate(pdf,1):p.get_pixmap(matrix=fitz.Matrix(1700/p.rect.width,1700/p.rect.width)).save(review/f'MDR-{i:03}.png')
print(json.dumps(summary))
