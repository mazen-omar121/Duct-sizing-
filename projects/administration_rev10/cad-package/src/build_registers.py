"""Create the Rev10 drawing-register Excel companion from the authoritative CSVs."""
from pathlib import Path
import csv,json
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter
BASE=Path(__file__).resolve().parents[1];OUT=BASE.parent
wb=Workbook();intro=wb.active;intro.title='Read_me'
readme=[
 ('ADMINISTRATION HVAC - REV10 DRAWING REGISTERS','07 October 2026 | engineering coordination issue'),
 ('Purpose','Editable CSV companion / sizing check for 21 detailed A1 sheets and four supplementary schematic sheets. Does not replace the original HAP workbook.'),
 ('Supply / return','Room duties retained: supply 3697.80 L/s; assumed return 3006.68 L/s. Pressure targets / outdoor-air balance remain unverified.'),
 ('Room connections','Clean-agent, kitchen and toilets have PAU supply and dedicated extract, without shared return; GF-07 bed/rest room retains return.'),
 ('Plant / fire','PAUs 3 x 50%, rotating pair of duty units; confirmed fire stops all HVAC fans and heaters. Fire-damper boundaries / ratings require approval.'),
 ('Routing','122 sections with aggregate-flow corridor headers. Rev10_route_changes records paths / flow / size updates; Rev10_size_changes compares final sizes directly against Rev09.'),
 ('Terminal changes','22 proposed face positions changed for separated routing / corridor service. Individual terminal duty, neck and face sizes are retained; approve ceiling / throws / noise before use.'),
 ('Geometry','Continuous service exteriors; replacement adapters; formed fitting throats; no unintended same-service polygon overlap or duplicated centerline. D-019 shows actual junctions.'),
 ('Sizing_check','Recalculated areas and velocities for revised clear sizes, not a fan ESP calculation. Horizontal centerline lengths and Colebrook straight-duct friction are recomputed. Vertical development and selected fitting K remain pending.'),
 ('Crossing_register','Separate-service footprints are numbered on plans and D-020. Calculated local maximum BODs are envelope bounds, not connected-network assigned elevations.'),
 ('Structural underside','User confirmed 4.000 m slab/roof underside. AFFL reference requires site verification; suspended finished ceiling remains unknown.'),
 ('Insulation','50 mm external indoor coordination assumption only; not a verified UAE average / specification. Outdoor / fire / vapour-barrier requirements need selected project specification.'),
 ('Height hold','CX-08: 750 + 800 mm bare ducts; 1950 mm required stack with assumed 50 mm insulation, 100 mm inter-duct gap and 100 mm structure allowance. Lowest insulated surface +2.050 m. Resolve shallow section / reroute against finished ceiling.'),
 ('Access / instruments','21 AD tags; 85 instrument requirements and I/O duties retained. Device cores checked on straight stations clear of crossing footprints. OEM dimensions, working access and safety settings require confirmation.'),
 ('I/O','65 AI / 22 AO / 141 DI / 44 DO and three COMM interfaces. Six explicit heater proof DI inputs retained; OEM safety chain independent of PLC demand.'),
 ('Holds','Actual pressure / transfer and outdoor air, levels / beams / access, fitting K / fan ESP / curves, selected equipment / diffuser performance, CAG / basement duties and final wiring remain open.'),
 ('How to edit','Update source JSON / CSV inputs together; regenerate drawings and workbook. Excel-only changes are not read by the generator. Original Rev08 and published Rev09 remain separate.'),
]
for row in readme:intro.append(row)
intro.column_dimensions['A'].width=28;intro.column_dimensions['B'].width=115
for row in intro:
 for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=11,color='192633')
 intro.row_dimensions[row[0].row].height=45 if row[0].row>1 else 30
for c in intro[1]:c.fill=PatternFill('solid',fgColor='123B52');c.font=Font(name='Calibri',size=14,bold=True,color='FFFFFF')
intro.freeze_panes='B2'

maps=[('Duct_sections.csv','Duct_sections'),('Transitions_and_takeoffs.csv','Transitions'),('Fittings.csv','Fittings'),('Terminals.csv','Terminals'),('Components.csv','Components'),('Device_locations.csv','Device_locations'),('Instrument_locations.csv','Probe_locations'),('Instruments.csv','Instruments'),('IO_points.csv','IO_points'),('Cause_and_effect.csv','Cause_effect'),('Rev08_corrections.csv','Rev08_corrections'),('Instrument_drawing_index.csv','Instrument_drawing_index'),('Rev10_route_changes.csv','Rev10_route_changes'),('Rev10_terminal_changes.csv','Rev10_terminal_changes'),('Rev10_size_changes.csv','Rev10_size_changes'),('Crossing_register.csv','Crossing_register'),('Formed_junctions.csv','Formed_junctions')]
for filename,title in maps:
 ws=wb.create_sheet(title)
 with (BASE/'schedules'/filename).open() as f:
  reader=csv.reader(f)
  for row in reader:
   vals=[]
   for val in row:
    # Preserve identifiers, room numbers and unresolved entries as text.
    vals.append(val)
   ws.append(vals)

ws=wb.create_sheet('Sizing_check')
ws.append(['Section','Air','Room','Flow L/s','Width mm','Height mm','Diameter mm','Area m2','Velocity m/s','Scheduled m/s','Difference m/s','Length m','Basis'])
from math import pi,ceil
rows=list(csv.DictReader((BASE/'schedules/Duct_sections.csv').open()))
for r in rows:
 q=float(r['flow_l_s']);w=float(r['width_mm']);d=float(r['diameter_mm']) if r['diameter_mm'] else None;h=float(r['height_mm']) if r['height_mm'] else None
 area=pi*(d/1000)**2/4 if d else w*h/1e6;v=q/1000/area;old=float(r['velocity_m_s'])
 ws.append([r['tag'],r['air_type'],r['room'],q,w,h,d,area,v,old,v-old,float(r['horizontal_length_m']),'Rev10 proposed clear size; retained room duty; recomputed velocity'])
 assert abs(v-old)<.00002,(r['tag'],v,old)
 for c in ws[ws.max_row][3:12]:c.number_format='0.000'

ws=wb.create_sheet('Comment_closure')
ws.append(['Comment','Rev10 action','Drawing / register','Status'])
for row in [
 ('Poor / broken connection walls','Render one complete service exterior after all fills; adapters replace original constant-width throat','D-001 to D-006 and actual geometry details D-019','Physical envelopes validated; actual delivered DXF exterior coverage checked'),
 ('Multiple same-service crossings / retraced runouts','Rebuild corridor collection/distribution as aggregate headers; separate meeting / telecom / operator / bed / clean-agent trees','Rev10_route_changes / D-019 / D-021','No unintended same-service polygon overlaps or duplicated centerlines'),
 ('Separate-service crossovers','Index physical footprints; dashed lower edges only; evaluate local stack against 4 m roof datum','D-020 / Crossing_register','BOD unassigned; resolve CX-08 height constraint before construction'),
 ('Dampers / heaters on joints and overlapping tags','Check reference glyph core dimension, straight wall station and crossing exclusions; use distinct leadered labels','Device_locations / D-004 to D-006 / D-019','Drawing footprint checked; OEM face-to-face and access pending'),
 ('Instrumentation completeness','Retain all 85 requirements and explicit fan / heater / basement / metering diagrams and sheet references','D-007 / D-012 to D-016 / Instrument_drawing_index','All 85 tags graphically referenced; signal duties unchanged'),
 ('Changes must be traceable','Document 22 terminal relocations and final section sizes directly against Rev09; recalculate section lengths / velocities / friction','D-021 / Rev10_terminal_changes / Rev10_size_changes / Sizing_check','Proposed positions / clear sizes require coordination and approval'),
]:ws.append(row)

for index,ws in enumerate(wb.worksheets[1:],1):
 ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
 ws.print_title_rows='1:1';ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='123B52');c.font=Font(name='Calibri',size=10,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=34
 for row in ws.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=10,color='192633')
  ws.row_dimensions[row[0].row].height=46 if ws.title in ['Instruments','IO_points','Cause_effect','Comment_closure','Rev08_corrections','Probe_locations'] else 32
 for j,col in enumerate(ws.iter_cols(),1):
  maximum=max(len(str(c.value or '')) for c in col);width=min(65,max(14,maximum*.8))
  ws.column_dimensions[get_column_letter(j)].width=width
 for row in ws.iter_rows(min_row=2):
  lines=max(sum(max(1,ceil(len(line)/(ws.column_dimensions[c.column_letter].width*.85))) for line in str(c.value or '').split('\n')) for c in row)
  ws.row_dimensions[row[0].row].height=max(30,lines*14+8)
 tab=Table(displayName='HVAC_'+str(index),ref=ws.dimensions);tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
 ws.add_table(tab)
 ws.sheet_properties.tabColor='228B22' if ws.title in ['Duct_sections','Transitions','Sizing_check'] else '703779' if ws.title in ['Instruments','IO_points','Cause_effect'] else '123B52'
wb.save(OUT/'Administration_HVAC_Drawing_Registers_Rev10.xlsx')
check=load_workbook(OUT/'Administration_HVAC_Drawing_Registers_Rev10.xlsx')
assert check['Duct_sections'].max_row==len(rows)+1
assert check['Instruments'].max_row==86
print(json.dumps({'Excel_sheets':len(wb.worksheets),'velocity_checks':len(rows),'registers':'Matched drawing CSVs'}))
