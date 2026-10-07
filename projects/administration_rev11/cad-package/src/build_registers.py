"""Create the Rev11 drawing-register Excel companion from the authoritative CSVs."""
from pathlib import Path
import csv,json
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter
BASE=Path(__file__).resolve().parents[1];OUT=BASE.parent
wb=Workbook();intro=wb.active;intro.title='Read_me'
readme=[
 ('ADMINISTRATION HVAC - REV11 ROOF OPTION','07 October 2026 | engineering coordination option, not for construction'),
 ('Drawings','22 detailed A1 sheets; four functional schematic sheets. D-019 actual junctions; D-020 height envelope; D-021 changes; D-022 proposed roof drops.'),
 ('User datums','Structural slab/roof underside +4.000 m; minimum ceiling / clear height +3.300 m. Verify AFFL and beam downstands on site.'),
 ('Strategy','Roof common mains with 15 independent room drops. USER SELECTED ROOF DISTRIBUTION / ROOM DROPS. Actual roof routes and penetrations require structural / architectural coordination.'),
 ('Indoor geometry','116 indoor paths; nine supply, six return and two extract connected room trees; no indoor service crossings. Proposed bare BOD +3.400 m, bare depth <=400 mm.'),
 ('Insulation','50 mm each face assumed indoors, not a verified UAE standard. Deepest envelope: insulated top +3.850 m, bottom +3.350 m, slab allowance 150 mm and ceiling gap 50 mm.'),
 ('Envelope limits','Only ducts checked: selected flanges, actuators, heater casings, supports, access and terminal plenums may require local redesign. Ceiling gap is not a general access / plenum allowance.'),
 ('Roof_main_basis','Retained common-main flow / clear-size basis, functional only. Actual roof plan / lengths / BOD / outdoor insulation / weatherproofing and losses TBC.'),
 ('Roof_risers','15 proposed source-plan drop centers. No structural penetration approval, selected sleeves / curbs / fire rating or actual roof vertical development.'),
 ('Revisions','13 additional terminal position proposals and 20 existing section size changes relative to Rev10. New perimeter headers are included in route changes.'),
 ('Sizing_check','All indoor velocities / plan centerline lengths / Colebrook straight friction recomputed. Not a complete fan ESP calculation; selected fitting K and roof / vertical lengths pending.'),
 ('Instrumentation','85 requirements and I/O duties retained. Common probes / heater / conditional humidifier move to proposed weather-protected roof plant. Five common probes shown functionally on D-012, not physically located on an unsurveyed roof.'),
 ('Pressure / OA','SA 3697.80 L/s; assumed RA 3006.68 L/s. OA reference 420.70 vs 691.12 implied makeup. Toilet net -18.69 L/s and kitchen zero do not prove +25 Pa.'),
 ('Access / protection','20 proposed indoor AD plan tags plus AD-HCOM-01 in roof plant schematic. OEM dimensions and working access, conditional fire boundaries and final wiring remain open.'),
 ('History','Rev08 / Rev09 / Rev10 originals retained. Rev10 stacked crossing CX-08 bottom +2.050 m fails new +3.300 m requirement and is superseded for this option.'),
 ('Editing','Edit source CSV/JSON together and regenerate. Excel-only edits are not imported. Do not rerun revision migrations after editing prepared inputs.'),
]
for row in readme:intro.append(row)
intro.column_dimensions['A'].width=28;intro.column_dimensions['B'].width=115
for row in intro:
 for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=11,color='192633')
 intro.row_dimensions[row[0].row].height=45 if row[0].row>1 else 30
for c in intro[1]:c.fill=PatternFill('solid',fgColor='123B52');c.font=Font(name='Calibri',size=14,bold=True,color='FFFFFF')
intro.freeze_panes='B2'

maps=[('Duct_sections.csv','Duct_sections'),('Transitions_and_takeoffs.csv','Transitions'),('Fittings.csv','Fittings'),('Terminals.csv','Terminals'),('Components.csv','Components'),('Device_locations.csv','Device_locations'),('Instrument_locations.csv','Probe_locations'),('Instruments.csv','Instruments'),('IO_points.csv','IO_points'),('Cause_and_effect.csv','Cause_effect'),('Rev08_corrections.csv','Rev08_corrections'),('Instrument_drawing_index.csv','Instrument_drawing_index'),('Rev11_route_changes.csv','Rev11_route_changes'),('Rev11_terminal_changes.csv','Rev11_terminal_changes'),('Rev11_size_changes.csv','Rev11_size_changes'),('Crossing_register.csv','Crossing_register'),('Formed_junctions.csv','Formed_junctions'),('Roof_main_basis.csv','Roof_main_basis'),('Roof_risers.csv','Roof_risers'),('Headroom_register.csv','Headroom_register')]
for filename,title in maps:
 ws=wb.create_sheet(title)
 with (BASE/'schedules'/filename).open() as f:
  reader=csv.reader(f)
  for row in reader:
   vals=[]
   for val in row:
    # Preserve identifiers, room numbers and unresolved entries as text.
    vals.append(val)
   ws.append(vals)

ws=wb.create_sheet('Sizing_check')
ws.append(['Section','Air','Room','Flow L/s','Width mm','Height mm','Diameter mm','Area m2','Velocity m/s','Scheduled m/s','Difference m/s','Length m','Basis'])
from math import pi,ceil
rows=list(csv.DictReader((BASE/'schedules/Duct_sections.csv').open()))
for r in rows:
 q=float(r['flow_l_s']);w=float(r['width_mm']);d=float(r['diameter_mm']) if r['diameter_mm'] else None;h=float(r['height_mm']) if r['height_mm'] else None
 area=pi*(d/1000)**2/4 if d else w*h/1e6;v=q/1000/area;old=float(r['velocity_m_s'])
 ws.append([r['tag'],r['air_type'],r['room'],q,w,h,d,area,v,old,v-old,float(r['horizontal_length_m']),'Rev11 proposed clear size; retained room duty; recomputed velocity'])
 assert abs(v-old)<.00002,(r['tag'],v,old)
 for c in ws[ws.max_row][3:12]:c.number_format='0.000'

ws=wb.create_sheet('Comment_closure')
ws.append(['Comment','Rev11 action','Drawing / register','Status'])
for row in [
 ('3.30 m clear height below 4.00 m slab','Limit indoor depth to 400 mm; propose bare BOD +3.400 m, roof common mains and separate room drops','D-020 / D-022 / Headroom_register','Duct envelope checked; Roof strategy selected; structural / OEM selections pending'),
 ('Multiple duct crossings','Perimeter return trees / independent room feeds replace ceiling common mains','D-001 to D-006 / D-019','Zero bare indoor crossings; separate-service 50 mm insulation envelopes do not overlap'),
 ('Connections / controls','Continuous exterior, replacement neck adapters, actual formed junctions and serial branch devices','D-019 / Device_locations / Formed_junctions','Connected room trees / station bodies checked; fabrication and access pending'),
 ('Instrumentation','Retain 85 requirements / I/O and relocate common-duct context to roof plant','D-012 to D-016 / Instrument_drawing_index','All tags referenced; roof physical mounting pending'),
 ('Roof extension','Separate functional common-main basis from indoor measured paths; schedule all 15 proposed risers','Roof_main_basis / Roof_risers / D-022','Actual roof routing, structure, weatherproofing, loads, levels and ESP not designed from survey'),
]:ws.append(row)

for index,ws in enumerate(wb.worksheets[1:],1):
 ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
 ws.print_title_rows='1:1';ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='123B52');c.font=Font(name='Calibri',size=10,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=34
 for row in ws.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=10,color='192633')
  ws.row_dimensions[row[0].row].height=46 if ws.title in ['Instruments','IO_points','Cause_effect','Comment_closure','Rev08_corrections','Probe_locations'] else 32
 for j,col in enumerate(ws.iter_cols(),1):
  maximum=max(len(str(c.value or '')) for c in col);width=min(65,max(14,maximum*.8))
  ws.column_dimensions[get_column_letter(j)].width=width
 for row in ws.iter_rows(min_row=2):
  lines=max(sum(max(1,ceil(len(line)/(ws.column_dimensions[c.column_letter].width*.85))) for line in str(c.value or '').split('\n')) for c in row)
  ws.row_dimensions[row[0].row].height=max(30,lines*14+8)
 tab=Table(displayName='HVAC_'+str(index),ref=ws.dimensions);tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)

 if ws.max_row>1:ws.add_table(tab)
 ws.sheet_properties.tabColor='228B22' if ws.title in ['Duct_sections','Transitions','Sizing_check'] else '703779' if ws.title in ['Instruments','IO_points','Cause_effect'] else '123B52'
wb.save(OUT/'Administration_HVAC_Drawing_Registers_Rev11.xlsx')
check=load_workbook(OUT/'Administration_HVAC_Drawing_Registers_Rev11.xlsx')
assert check['Duct_sections'].max_row==len(rows)+1
assert check['Instruments'].max_row==86
print(json.dumps({'Excel_sheets':len(wb.worksheets),'velocity_checks':len(rows),'registers':'Matched drawing CSVs'}))
